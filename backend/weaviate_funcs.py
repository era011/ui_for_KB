from weaviate.classes.query import Filter
from backend.db import *
from functools import lru_cache
import os
import io
import hashlib
from langchain_core.documents import Document
from const import *
from weaviate.classes.config import Property, DataType, Configure
import datetime
from typing import List, Tuple
from backend.weaviate_client import connect_to_local
import re
from dotenv import load_dotenv
from const import *
from langchain_text_splitters import RecursiveCharacterTextSplitter

load_dotenv()


@lru_cache(maxsize=1)
def get_client():
    return connect_to_local(
        host=os.getenv("SERVER"),
        port=int(os.getenv("WEAVIATE_PORT")),
        grpc_port=50051,
        headers={"X-OpenAI-Api-Key": os.getenv("OPENAI_API_KEY")}
    )


# def ensure_schema(client, collection_name: str):
#     existing = client.collections.list_all()
#     print(type(existing))
#     if collection_name not in existing:
#         client.collections.create(
#             name=collection_name,
#             vectorizer_config=Configure.Vectorizer.text2vec_openai(
#             model="text-embedding-3-large"
#             ),
#             properties=[
#                 Property(name="content", data_type=DataType.TEXT,vectorize=True),
#                 Property(name="name", data_type=DataType.TEXT,vectorize=False),
#                 Property(name="id_doc", data_type=DataType.TEXT,vectorize=False),
#                 Property(name="added_date_to_weaviate", data_type=DataType.DATE,vectorize=False),
#             ]
#         )
#         print(f"✅ Коллекция {collection_name} создана")
#     else:
#         print(f"ℹ️ Коллекция {collection_name} уже существует")

def ensure_schema(client, collection_name: str):
    try:
        client.collections.get(collection_name)
        print(f"ℹ️ Коллекция {collection_name} уже существует")
        return
    except Exception:
        # коллекции нет -> создаём
        client.collections.create(
            name=collection_name,
            vectorizer_config=Configure.Vectorizer.text2vec_openai(
                model="text-embedding-3-large"
            ),
            properties=[
                Property(name="content", data_type=DataType.TEXT, vectorize=True),
                Property(name="name", data_type=DataType.TEXT, vectorize=False),
                Property(name="id_doc", data_type=DataType.TEXT, vectorize=False),
                Property(name="added_date_to_weaviate", data_type=DataType.DATE, vectorize=False),
                Property(name="chunk_index", data_type=DataType.INT, vectorize=False),
            ]
        )
        print(f"✅ Коллекция {collection_name} создана")

def delete_document_weaviate(client, id_doc: str, collection_name: str = COLLECTION_NAME):
    try:
        collection = client.collections.get(collection_name)
        result = collection.query.fetch_objects(
            filters=Filter.by_property("id_doc").equal(id_doc),
            limit=1000
        )
        for obj in result.objects:
            collection.data.delete_by_id(obj.uuid)
        print(f"🗑 Документ {id_doc} удалён из Weaviate")
    except Exception as e:
        raise ValueError(f"Ошибка при удалении из Weaviate: {e}")
    
def get_chunks_weaviate(client, doc_id: str, limit: int = 1000):
    try:
        collection = client.collections.get(COLLECTION_NAME)
        result = collection.query.fetch_objects(
            filters=Filter.by_property("id_doc").equal(doc_id),
            limit=limit
        )
        chunks = []
        for obj in result.objects:
            chunks.append(obj.properties.get("content", ""))
        if result.objects:
            last = result.objects[-1]
            chunks.append(str({
                "chunk_index": last.properties.get("chunk_index", ""),
                "name": last.properties.get("name", ""),
                "id_doc": last.properties.get("id_doc", ""),
                "added_date_to_weaviate": last.properties.get("added_date_to_weaviate", ""),
                }))
        return chunks
    except Exception as e:
        raise ValueError(f"Ошибка при получении чанков из Weaviate: {e}")


def document_not_exists_pg(id_doc: str) -> bool:
    conn = None
    try:
        conn = get_connection()
        cur = conn.cursor()
        cur.execute(f"SELECT 1 FROM {PG_TABLE_NAME} WHERE id_doc = %s LIMIT 1;", (id_doc,))
        exists = cur.fetchone() is not None
        cur.close()
        return not exists
    except Exception as e:
        raise ValueError(f"Ошибка при проверке документа в Postgres: {e}")
    finally:
        if conn:
            release_connection(conn)

def doc_processing(text:str, hash:str, name:str):
    try:
        docs = Document(page_content=text, metadata={
            "name": name,
            "id_doc": hash,
            "added_date_to_weaviate": datetime.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        })
        return docs
    except Exception as e:
        raise ValueError(f"Произошла ошибка при обработке документа: {e}")


# def split_by_sections(doc: Document) -> List[Document]:
#     """
#     Делит документ на чанки по двум и более пустым строкам.
#     Возвращает список langchain_core.documents.Document
#     """
#     if not doc.page_content or not doc.page_content.strip():
#         return []
#     text = doc.page_content.replace("\r\n", "\n").replace("\r", "\n")
#     sections = re.split(r"\n\s*\n\s*\n+", text)
#     chunks: List[Document] = []
#     for idx, section in enumerate(sections):
#         section = section.strip()
#         if not section:
#             continue
#         chunk = Document(
#             page_content=section,
#             metadata={
#                 **(doc.metadata or {}),
#                 "chunk_index": idx
#             }
#         )
#         chunks.append(chunk)
#     return (chunks, len(chunks))

def split_by_sections_then_textsplitter(doc: Document, chunk_size: int = 1000, chunk_overlap: int = 300,) -> Tuple[List[Document], int]:
    """
    1) Делит документ на секции по 2+ пустым строкам
    2) Каждую секцию дробит RecursiveCharacterTextSplitter'ом
    Возвращает (chunks, len(chunks))
    """
    if not doc.page_content or not doc.page_content.strip():
        return ([], 0)
    text = doc.page_content.replace("\r\n", "\n").replace("\r", "\n")
    sections = re.split(r"\n\s*\n\s*\n+", text)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],  # можно менять под твои документы
    )
    out: List[Document] = []
    global_idx = 0
    for section_idx, section in enumerate(sections):
        section = section.strip()
        if not section:
            continue

        # режем секцию на под-чанки
        sub_texts = splitter.split_text(section)

        for sub_idx, sub_text in enumerate(sub_texts):
            out.append(
                Document(
                    page_content=sub_text,
                    metadata={
                        **(doc.metadata or {}),
                        "section_index": section_idx,
                        "subchunk_index": sub_idx,
                        "chunk_index": global_idx,  # глобальный индекс для Weaviate
                    },
                )
            )
            global_idx += 1
    return (out, len(out))

def upload_chunks(client, collection_name: str, chunks: List[Document]):
    collection = client.collections.get(collection_name)
    for idx, chunk in enumerate(chunks):
        collection.data.insert(
            properties={
                "content": chunk.page_content,
                "name": chunk.metadata.get("name", ""),
                "id_doc": chunk.metadata.get("id_doc", ""),
                "added_date_to_weaviate": chunk.metadata.get("added_date_to_weaviate", ""),
                "chunk_index": chunk.metadata.get("chunk_index", 0),
            },
        )

SUPPORTED_EXTENSIONS = (".txt", ".pdf", ".docx")


def _decode_txt_bytes(bytes_data: bytes) -> str:
    try:
        return bytes_data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return bytes_data.decode("cp1251")
        except UnicodeDecodeError as e:
            raise ValueError("Не удалось декодировать файл. Нужен UTF-8 (или cp1251).") from e


def _extract_text_from_pdf(bytes_data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("Для загрузки PDF нужен пакет pypdf (pip install pypdf).") from e
    try:
        reader = PdfReader(io.BytesIO(bytes_data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages).strip()
    except Exception as e:
        raise ValueError(f"Не удалось прочитать PDF-файл: {e}") from e


def _extract_text_from_docx(bytes_data: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as e:
        raise ValueError("Для загрузки DOCX нужен пакет python-docx (pip install python-docx).") from e
    try:
        docx_doc = DocxDocument(io.BytesIO(bytes_data))
        parts = [p.text for p in docx_doc.paragraphs if p.text.strip()]
        for table in docx_doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n\n".join(parts).strip()
    except Exception as e:
        raise ValueError(f"Не удалось прочитать DOCX-файл: {e}") from e


def extract_text(name: str, bytes_data: bytes) -> str:
    ext = os.path.splitext(name)[1].lower()
    if ext == ".txt":
        return _decode_txt_bytes(bytes_data)
    elif ext == ".pdf":
        return _extract_text_from_pdf(bytes_data)
    elif ext == ".docx":
        return _extract_text_from_docx(bytes_data)
    else:
        raise ValueError(
            f"Неподдерживаемый формат файла: {ext or 'неизвестен'}. "
            f"Поддерживаются: {', '.join(SUPPORTED_EXTENSIONS)}"
        )


def add_text_document_to_weaviate(uploaded_file):
    """
    Добавляет документ (.txt, .pdf, .docx) в Postgres + Weaviate.
    expected: uploaded_file.read() -> bytes, uploaded_file.name -> str
    Возвращает id_doc добавленного документа.
    """
    try:
        client = get_client()
        bytes_data = uploaded_file.read()
        name = getattr(uploaded_file, "name", "uploaded.txt")
        text = extract_text(name, bytes_data)
        if not text.strip():
            raise ValueError(f"Не удалось извлечь текст из документа {name}")
        doc_hash = hashlib.md5(text.encode("utf-8")).hexdigest()
        ensure_schema(client, COLLECTION_NAME)
        if not document_not_exists_pg(doc_hash):
            raise ValueError(f"Документ {name} уже существует")
        else:
            try:
                doc = doc_processing(text, doc_hash, name)
                chunks, chunks_count = split_by_sections_then_textsplitter(doc)
                add_document_to_postgres(doc, chunks_count)
                upload_chunks(client, COLLECTION_NAME, chunks)
            except Exception as e:
                delete_document_postgres(doc_hash)
                raise ValueError(f"Произошла ошибка при добавлении документа в Weaviate: {e}")
        return doc_hash

    except Exception as e:
        raise ValueError(f"Произошла ошибка при добавлении документа: {e}")